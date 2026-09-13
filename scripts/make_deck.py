"""Builds the four-slide Quorly deck.

Same shape as the Renesis deck it is modelled on: white ground, Anton set huge,
a dry subtitle, a meme doing the emotional work, and a punchline bottom right.

The memes are generated with our own captions through memegen.link rather than
pulled off an image search, so the joke is about this product rather than a
stock gag with a caption bolted on. They live in assets/deck/ so the deck
rebuilds without a network.

    python scripts/make_deck.py            # writes Quorly_Deck.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

W, H = Inches(13.333), Inches(7.5)  # 16:9
MEMES = Path(__file__).resolve().parent.parent / "assets" / "deck"

INK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
BLUE = RGBColor(0x08, 0x51, 0xBF)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x1E, 0x7A, 0x4E)
PAPER = RGBColor(0xF7, 0xF6, 0xF2)
LINE = RGBColor(0xE2, 0xE0, 0xD9)

HEAD = "Anton"
BODY = "Helvetica Neue"


def deck():
    p = Presentation()
    p.slide_width, p.slide_height = W, H
    return p


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    return s


def text(slide, x, y, w, h, runs, size=18, font=BODY, color=INK,
         align=PP_ALIGN.LEFT, bold=False, spacing=1.0, anchor=MSO_ANCHOR.TOP):
    """runs is a string, or a list of (text, {overrides}) for mixed styling."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    if isinstance(runs, str):
        runs = [(runs, {})]

    para = tf.paragraphs[0]
    para.alignment = align
    para.line_spacing = spacing
    for content, over in runs:
        r = para.add_run()
        r.text = content
        f = r.font
        f.name = over.get("font", font)
        f.size = Pt(over.get("size", size))
        f.bold = over.get("bold", bold)
        f.color.rgb = over.get("color", color)
    return box


def card(slide, x, y, w, h, fill=PAPER, line=LINE):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    return shp


def meme(slide, name, x, y, height):
    """Places a meme at a fixed height, centred on x."""
    pic = slide.shapes.add_picture(str(MEMES / f"{name}.png"), x, y, height=height)
    pic.left = int(x - pic.width / 2)
    return pic


def headline(slide, big, sub=None, size=44):
    # Sized to sit on one line even if Anton is missing and something wider
    # substitutes: a headline that wraps lands on the subtitle, and a deck that
    # only lays out on the author's laptop is worse than a plainer one.
    text(slide, Inches(0.4), Inches(0.42), Inches(12.53), Inches(1.0), big,
         size=size, font=HEAD, align=PP_ALIGN.CENTER)
    if sub:
        text(slide, Inches(1.2), Inches(1.42), Inches(10.9), Inches(0.5), sub,
             size=17, color=MUTED, align=PP_ALIGN.CENTER)


def punchline(slide, s):
    text(slide, Inches(6.5), Inches(6.85), Inches(5.9), Inches(0.35), s,
         size=12, color=MUTED, align=PP_ALIGN.RIGHT)


# ─────────────────────────────────────────────────────────── 1. the problem

def slide_problem(prs):
    s = blank(prs)
    headline(s, "SOMEBODY HAS TO CLICK APPROVE",
             "and right now a stolen laptop can do it just as well as they can")

    meme(s, "fine", Inches(3.5), Inches(2.15), Inches(3.6))

    card(s, Inches(6.9), Inches(2.15), Inches(5.55), Inches(3.6),
         fill=RGBColor(0xFD, 0xF3, 0xF2), line=RGBColor(0xF0, 0xD5, 0xD2))
    text(s, Inches(7.35), Inches(2.5), Inches(4.7), Inches(0.35),
         "WHAT THAT CLICK IS WORTH", size=13, font=HEAD, color=RED)
    text(s, Inches(7.35), Inches(3.0), Inches(4.7), Inches(2.0),
         "Whatever the treasury holds.\n\n"
         "The invoice waits four days in an inbox nobody opens. Then somebody "
         "clicks a button that trusts a session cookie — and a click is the one "
         "thing that cannot tell you who clicked.",
         size=15, spacing=1.3)
    text(s, Inches(7.35), Inches(5.15), Inches(4.7), Inches(0.4),
         "bill.com takes days. Stealing the cookie takes seconds.",
         size=14, bold=True, color=RED)

    punchline(s, "(the dog is our approvals process.)")


# ────────────────────────────────────────────────────────── 2. architecture

def slide_architecture(prs):
    s = blank(prs)
    headline(s, "FOUR CHECKPOINTS, ONE PATH",
             "each one somewhere a stolen session cannot reach")

    meme(s, "drake", Inches(2.35), Inches(2.05), Inches(3.9))

    steps = [
        ("01", "SLACK", "Contractor DMs a PDF.\nClaude reads the amount.", INK),
        ("02", "POLICY", "Routed on amount and role.\nOnly real approvers pinged.", INK),
        ("03", "YOUR FACE", "Above the threshold, a live\nWorld ID Selfie Check.", BLUE),
        ("04", "THE QUORUM", "2-of-3 keys sign in an\nenclave. Then it pays.", GREEN),
    ]
    x, y = Inches(4.75), Inches(2.05)
    for i, (num, title, body, colour) in enumerate(steps):
        col, row = divmod(i, 2)
        cx = x + Inches(3.95) * col
        cy = y + Inches(2.0) * row
        card(s, cx, cy, Inches(3.7), Inches(1.8))
        text(s, cx + Inches(0.3), cy + Inches(0.22), Inches(1.0), Inches(0.4), num,
             size=22, font=HEAD, color=colour)
        text(s, cx + Inches(1.0), cy + Inches(0.26), Inches(2.4), Inches(0.35), title,
             size=15, font=HEAD)
        text(s, cx + Inches(0.3), cy + Inches(0.82), Inches(3.1), Inches(0.9), body,
             size=12, color=MUTED, spacing=1.2)

    card(s, Inches(0.75), Inches(6.15), Inches(11.8), Inches(0.62),
         fill=RGBColor(0xF2, 0xF7, 0xFF), line=RGBColor(0xD5, 0xE3, 0xFA))
    text(s, Inches(1.1), Inches(6.31), Inches(11.1), Inches(0.4),
         [("The bit that matters:  ", {"bold": True}),
          ('"two approvals required" is not a flag in our database. It is the '
           "wallet's owner. We could not skip it if we wanted to.", {})],
         size=13)

    punchline(s, "Go · Postgres queue · Next.js · Base Sepolia")


# ──────────────────────────────────────────────────────────── 3. why quorly

def slide_why(prs):
    s = blank(prs)
    headline(s, "WHY QUORLY", "built in a hackathon, wired to real infrastructure", 50)

    for i, (title, body) in enumerate([
        ("Approvals live in Slack", "Where the work already happens. No portal, no new login."),
        ("A face, not a session", "Selfie Check, bound to one invoice, expires in minutes."),
        ("A treasury nobody owns alone", "2-of-3 key quorum, spend policy enforced in an enclave."),
        ("Paid in about a minute", "Not four days. The hash lands on the invoice."),
    ]):
        y = Inches(2.2) + Inches(0.92) * i
        text(s, Inches(0.85), y, Inches(5.4), Inches(0.3), title, size=16, bold=True)
        text(s, Inches(0.85), y + Inches(0.32), Inches(5.4), Inches(0.4), body,
             size=13, color=MUTED)

    meme(s, "success", Inches(8.0), Inches(2.1), Inches(2.35))

    card(s, Inches(6.6), Inches(4.75), Inches(2.85), Inches(1.9),
         fill=RGBColor(0xF1, 0xF8, 0xF4), line=RGBColor(0xCF, 0xE6, 0xDA))
    text(s, Inches(6.9), Inches(4.98), Inches(2.3), Inches(0.3),
         "ALREADY ONCHAIN", size=11, font=HEAD, color=GREEN)
    text(s, Inches(6.9), Inches(5.32), Inches(2.35), Inches(1.2),
         "$2,000 released after a\nlive Selfie Check.\n2-of-3 signed, Base Sepolia.",
         size=12, spacing=1.25)

    card(s, Inches(9.65), Inches(4.75), Inches(2.9), Inches(1.9))
    text(s, Inches(9.95), Inches(4.98), Inches(2.3), Inches(0.3),
         "NEXT", size=11, font=HEAD, color=MUTED)
    text(s, Inches(9.95), Inches(5.32), Inches(2.4), Inches(1.2),
         "Settle in local currency —\nrupees, pesos, naira land,\nnot a token to go and sell.",
         size=12, spacing=1.25)

    punchline(s, "quorly.xyz")


# ─────────────────────────────────────────────────────────────── 4. thanks

def slide_thanks(prs):
    s = blank(prs)

    text(s, Inches(0.6), Inches(1.5), Inches(12.13), Inches(1.4), "THANK YOU",
         size=80, font=HEAD, align=PP_ALIGN.CENTER)
    text(s, Inches(2.4), Inches(2.75), Inches(8.5), Inches(0.5),
         "Money moves when a live human says so.",
         size=20, color=MUTED, align=PP_ALIGN.CENTER)

    meme(s, "cheers", Inches(6.67), Inches(3.35), Inches(2.5))

    card(s, Inches(4.15), Inches(6.05), Inches(5.0), Inches(0.95))
    text(s, Inches(4.35), Inches(6.24), Inches(4.6), Inches(0.3),
         "quorly.xyz", size=17, bold=True, align=PP_ALIGN.CENTER)
    text(s, Inches(4.35), Inches(6.58), Inches(4.6), Inches(0.3),
         "github.com/Nishu0/quorly  ·  itsnishu", size=12, color=MUTED,
         align=PP_ALIGN.CENTER)

    punchline(s, "(the invoice was, in fact, paid.)")


def main():
    prs = deck()
    slide_problem(prs)
    slide_architecture(prs)
    slide_why(prs)
    slide_thanks(prs)
    prs.save("Quorly_Deck.pptx")
    print("wrote Quorly_Deck.pptx")


if __name__ == "__main__":
    main()
